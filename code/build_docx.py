# -*- coding: utf-8 -*-
"""Build the manuscript and supplementary Word files."""
import os
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

FONT = "Times New Roman"


def _style(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(doc, text, size=12, bold=False, italic=False, align="justify",
         space_before=0, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
                   "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return p


def add_table(doc, df, caption, max_rows=40, fontsize=9):
    para(doc, caption, size=10, bold=False, align="left", space_before=10, space_after=4)
    d = df.head(max_rows).fillna("")
    t = doc.add_table(rows=1, cols=len(d.columns))
    t.style = "Table Grid"
    for j, c in enumerate(d.columns):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(c).replace("_", " "))
        r.bold = True
        r.font.size = Pt(fontsize)
        r.font.name = FONT
    for _, row in d.iterrows():
        cells = t.add_row().cells
        for j, v in enumerate(row):
            if isinstance(v, float):
                v = ("%.3f" % v).rstrip("0").rstrip(".") if abs(v) < 1000 else "{:,.0f}".format(v)
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(str(v))
            r.font.size = Pt(fontsize)
            r.font.name = FONT
    doc.add_paragraph()


def add_figure(doc, path, caption, width=6.3):
    doc.add_picture(path, width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, caption, size=10, align="justify", space_before=4, space_after=12)


def word_count(blocks):
    return sum(len(t.split()) for s, t in blocks if s in ("P", "ABS"))
