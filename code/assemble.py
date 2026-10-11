# -*- coding: utf-8 -*-
"""Assemble main manuscript and supplementary Word files."""
import json
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from build_docx import _style, para, add_table, add_figure
import manuscript_final as M

BIB = M.BIB

MAIN_FIGS_NEW = [
    ("screen_architecture.png",
     "Figure 1. Architecture of the acellular cementum CRISPR knockout screen. The design "
     "proceeds from a two-tier target space (curated cementum module plus the complete "
     "kinome) through genome-verified guide enumeration, three on-target activity models "
     "and exhaustive off-target enumeration, to a synthesis-ready oligonucleotide pool. "
     "The pool feeds a planned pooled knockout screen in human periodontal ligament "
     "progenitors read out by fluorescence-activated sorting on a mineralisation reporter, "
     "with four analysis layers stacked on the resulting counts. No experimental arm of "
     "this architecture has been executed; every quantity reported in this work is a "
     "computational property of the design."),
    ("pipeline_funnel.png",
     "Figure 2. Guide design funnel and the composition of the final target space. "
     "(a) Sequential attrition from 139,144 enumerated candidate protospacers to the "
     "7,751 guides retained in the main design pass, an attrition of 94.4%; gene counts "
     "at each stage are given in parentheses. (b) Guides per submodule in the final "
     "8,127-guide targeting library, coloured by tier; the kinome accounts for 55.5% of "
     "the library."),
    ("offtarget_and_library.png",
     "Figure 3. Exhaustive off-target enumeration across the human genome. Every selected "
     "protospacer was expanded into all sequences within three substitutions and matched "
     "against every NGG, NAG and NGA-adjacent 20-mer in GRCh38. (a) Distribution of "
     "recovered sites by mismatch count; within the principal three-mismatch pass, 91.3% "
     "of the 514,765 sites lie at exactly three mismatches. (b) Aggregate CFD specificity "
     "does not separate the screened pool from the selected library (medians 0.265 and "
     "0.277), which is why selection used the maximum coding-exon CFD instead. (c) Guides "
     "per gene in the final library."),
    ("power_analysis.png",
     "Figure 4. Statistical power of the planned sorted screen under a negative-binomial "
     "count model. (a, b) Power across effect size, replicate number and library "
     "representation, computed on the coverage-resolved grid (Supplementary Table S2); "
     "replicate number dominates and representation above 200 cells per guide adds "
     "essentially nothing. (c) Sensitivity of power to the biological overdispersion "
     "parameter, computed on a separate grid spanning 6 to 12 guides per gene at four "
     "replicates; the highlighted row is the value measured from published replicate "
     "counts. Power figures quoted in the text come from the coverage-resolved grid in "
     "(a, b)."),
    ("scale_feasibility.png",
     "Figure 5. Cell number, sorting time and sequencing scale for the focused design "
     "against a genome-wide equivalent at the same decile gating. The focused screen "
     "requires 350 million cells and a 22-minute sort per replicate; the genome-wide "
     "equivalent requires 2.81 billion cells and a 2.9-hour sort, which is the "
     "quantitative basis for restricting the library to a focused target space."),
    ("network_module.png",
     "Figure 6. Graph neural network propagation over the STRING interactome. A two-layer "
     "mean-aggregator GraphSAGE model was trained on 817 nodes and 5,824 high-confidence "
     "edges with GTEx median expression across 54 tissues as node features. Held-out "
     "AUROC for recovery of cementum module membership was 0.805 ± 0.088 against 0.704 ± "
     "0.060 for logistic regression on the same features without graph structure (paired "
     "Wilcoxon over 20 splits, p = 2 × 10⁻⁴); a degree-only control reached 0.387."),
    ("readout_specificity.png",
     "Figure 7. The cementogenic and osteogenic programmes are not computationally "
     "separable. (a) Gene Ontology overlap between the tooth and bone programmes "
     "(Jaccard 0.054). (b) Held-out AUROC for discriminating tooth-only from bone-only "
     "genes reached only 0.582 ± 0.084, against 0.805 for module recovery on the same "
     "graph and features. (c) Coverage of the two programmes by the curated module before "
     "rebalancing, showing the 2.45-to-1 lean towards bone."),
    ("rebalancing.png",
     "Figure 8. Module rebalancing and the composition of the synthesis-ready pool. "
     "(a) Addition of 39 tooth-programme genes moved the bone-to-tooth ratio from 2.45:1 "
     "to 0.93:1. (b) Coverage of the Gene Ontology tooth programme rose from 44% to 99% "
     "for 376 additional guides. (c) The final pool of 9,631 oligonucleotides, of which "
     "15.6% are controls."),
]

MAIN_TABLES_NEW = [
    ("table1_pipeline_funnel.csv",
     "Table 1. Guide design funnel. Guides and genes remaining after each filtering stage "
     "of the main design pass."),
    ("table2_target_space.csv",
     "Table 2. Final target space by tier and submodule, with guides per gene."),
    ("table4_offtarget_passes.csv",
     "Table 3. Off-target screening passes. Guides screened, mismatch ceiling, wall-clock "
     "time and sites catalogued for each of the five genome-wide passes."),
    ("table7_library_composition.csv",
     "Table 4. Composition of the synthesis-ready oligonucleotide pool by guide class."),
]

SUP_FIGS_NEW = [
    ("library_composition.png",
     "Supplementary Figure S1. Composition of the final library. (a) Oligonucleotides by "
     "class. (b) Guides per target gene. (c) Aggregate CFD specificity of targeting guides "
     "against the two control classes, confirming that controls are neither systematically "
     "cleaner nor dirtier than the guides they calibrate."),
    ("flanks_and_budget.png",
     "Supplementary Figure S2. Oligonucleotide architecture and cost. (a, b) The 3′ "
     "homology flank is the sgRNA scaffold itself, so the original Hsu 2013 and F+E Chen "
     "2013 constant regions give different oligonucleotide strings; after off-target "
     "filtering the selected guide set is identical under either scaffold. (c) Incremental "
     "cost of the parallel osteoblast arm."),
]

SUP_TABLES_NEW = [
    ("table3_ontarget_models.csv",
     "Supplementary Table S1. Rank concordance (Spearman ρ) between the three on-target "
     "activity models applied to the selected guides."),
    ("tableS2_power_grid.csv",
     "Supplementary Table S2. Monte Carlo power grid across effect size, guides per gene, "
     "replicate number and library representation. Power is the fraction of simulated "
     "genes called at a 5% false discovery rate; the final column averages the three "
     "coverage levels. Values quoted in the Results are taken from this grid."),
    ("tableS3_scale_and_depth_T.csv",
     "Supplementary Table S3. Cell number, sorting time, sample count and sequencing "
     "depth for the focused and genome-wide designs."),
    ("table8_budget.csv",
     "Supplementary Table S4. Itemised budget for the core screen and the incremental "
     "cost of the parallel osteoblast arm."),
    ("table5_initial_vs_final.csv",
     "Supplementary Table S5. The screen as originally specified against the design "
     "delivered, with the analysis that motivated each change."),
]

MAIN_FIGS = [
    ("screen_architecture.png",
     "Figure 1. Architecture of the acellular cementum CRISPR knockout screen. The design "
     "proceeds from a two-tier target space (curated cementum module plus the complete "
     "kinome) through genome-verified guide enumeration, three on-target activity models "
     "and exhaustive off-target enumeration, to a synthesis-ready oligonucleotide pool. "
     "The pool feeds a planned pooled knockout screen in human periodontal ligament "
     "progenitors read out by fluorescence-activated sorting on a mineralisation reporter, "
     "with four analysis layers stacked on the resulting counts. No experimental arm of "
     "this architecture has been executed; all quantities reported in this work are "
     "computational properties of the design."),
    ("pipeline_funnel.png",
     "Figure 2. Guide design funnel and the composition of the final target space. "
     "(a) Sequential attrition from 139,144 enumerated candidate protospacers to the "
     "7,751 guides retained in the main design pass, an attrition of 94.4%; gene counts "
     "at each stage are given in parentheses. (b) Guides per submodule in the final "
     "8,127-guide targeting library, coloured by tier; the kinome accounts for 55.5% of "
     "the library."),
    ("offtarget_and_library.png",
     "Figure 3. Exhaustive off-target enumeration across the human genome. Every selected "
     "protospacer was expanded into all sequences within three substitutions and matched "
     "against every NGG, NAG and NGA-adjacent 20-mer in GRCh38. (a) Distribution of the "
     "794,259 catalogued sites by mismatch count; within the principal three-mismatch "
     "pass, 91.3% of the 514,765 recovered sites lie at exactly three mismatches. "
     "(b) Aggregate CFD specificity does not separate the screened pool from the selected "
     "library (medians 0.265 and 0.277), which is why selection was performed on the "
     "maximum coding-exon CFD instead."),
    ("network_module.png",
     "Figure 4. Graph neural network propagation over the STRING interactome. A two-layer "
     "mean-aggregator GraphSAGE model was trained on 817 nodes and 5,824 high-confidence "
     "edges with GTEx median expression across 54 tissues as node features. Held-out "
     "AUROC for recovery of cementum module membership was 0.805 ± 0.088 against 0.704 ± "
     "0.060 for logistic regression on the same features without graph structure (paired "
     "Wilcoxon over 20 splits, p = 2 × 10⁻⁴); a degree-only control reached 0.387."),
    ("readout_specificity.png",
     "Figure 5. The cementogenic and osteogenic programmes are not computationally "
     "separable. (a) Gene Ontology overlap between the tooth and bone programmes "
     "(Jaccard 0.054). (b) Held-out AUROC for discriminating tooth-only from bone-only "
     "genes reached only 0.582 ± 0.084, against 0.805 for module recovery on the same "
     "graph and features. (c) Coverage of the two programmes by the curated module before "
     "rebalancing, showing the 2.45-to-1 lean towards bone."),
    ("rebalancing.png",
     "Figure 6. Module rebalancing and the composition of the synthesis-ready pool. "
     "(a) Addition of 39 tooth-programme genes moved the bone-to-tooth ratio from 2.45:1 "
     "to 0.93:1. (b) Coverage of the Gene Ontology tooth programme rose from 44% to 99% "
     "for 376 additional guides. (c) The final pool of 9,631 oligonucleotides, of which "
     "15.6% are controls."),
]

MAIN_TABLES = [
    ("table1_pipeline_funnel.csv",
     "Table 1. Guide design funnel. Guides and genes remaining after each filtering stage "
     "of the main design pass."),
    ("table2_target_space.csv",
     "Table 2. Final target space by tier and submodule, with guides per gene."),
    ("table4_offtarget_passes.csv",
     "Table 3. Off-target screening passes. Guides screened, mismatch ceiling, wall-clock "
     "time and sites catalogued for each of the five genome-wide passes."),
    ("table7_library_composition.csv",
     "Table 4. Composition of the synthesis-ready oligonucleotide pool by guide class."),
]

SUP_FIGS = [
    ("power_analysis.png",
     "Supplementary Figure S1. Statistical power of the planned sorted screen under a "
     "negative-binomial count model with overdispersion measured from GSE226977 "
     "(φ = 0.299). Power is shown as a function of effect size, replicate number and "
     "library representation; replicate number dominates, and coverage above 200× adds "
     "essentially nothing."),
    ("scale_feasibility.png",
     "Supplementary Figure S2. Cell number, sorting time and sequencing scale for the "
     "focused design against a genome-wide equivalent at the same decile gating. The "
     "focused screen requires 350 million cells and a 22-minute sort per replicate; the "
     "genome-wide equivalent requires 2.81 billion cells and a 2.9-hour sort."),
    ("flanks_and_budget.png",
     "Supplementary Figure S3. Oligonucleotide architecture and cost. The 3′ homology "
     "flank is the sgRNA scaffold itself, so the original Hsu 2013 and F+E Chen 2013 "
     "constant regions give different oligonucleotide strings. Both variants were emitted "
     "and verified; the budget panel shows the incremental cost of the parallel "
     "osteoblast arm."),
    ("library_composition.png",
     "Supplementary Figure S4. Composition of the final library by guide class, by "
     "submodule and by guides per gene, with the distribution of exon spanning across "
     "targeted genes."),
]

SUP_TABLES = [
    ("table3_ontarget_models.csv",
     "Supplementary Table S1. Rank concordance (Spearman ρ) between the three on-target "
     "activity models applied to the selected guides."),
    ("tableS2_power_grid.csv",
     "Supplementary Table S2. Monte Carlo power grid across effect size, guides per gene, "
     "replicate number and library representation. Power is the fraction of simulated "
     "genes called at a 5% false discovery rate; the final column averages the three "
     "coverage levels."),
    ("tableS3_scale_and_depth_T.csv",
     "Supplementary Table S3. Cell number, sorting time, sample count and sequencing "
     "depth for the focused and genome-wide designs."),
    ("table8_budget.csv",
     "Supplementary Table S4. Itemised budget for the core screen and the incremental "
     "cost of the parallel osteoblast arm."),
    ("table5_initial_vs_final.csv",
     "Supplementary Table S5. The screen as originally specified against the design "
     "delivered, with the analysis that motivated each change."),
]


def build_main(path="Cementum_CRISPR_manuscript.docx"):
    doc = Document()
    _style(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
    para(doc, M.TITLE, size=14, bold=True, align="center", space_after=14)
    body = 0
    for style, text in M.BLOCKS:
        if style == "H1":
            para(doc, text, size=12, bold=True, align="left", space_before=12, space_after=6)
        elif style == "H2":
            para(doc, text, size=12, bold=False, italic=True, align="left",
                 space_before=8, space_after=4)
        elif style == "KW":
            para(doc, text, size=12, italic=True)
        elif style in ("P", "ABS"):
            para(doc, text, size=12)
            body += len(text.split())
    # References
    para(doc, M.REFERENCES_HEADER, size=12, bold=True, align="left",
         space_before=14, space_after=6)
    for r in BIB:
        para(doc, r, size=11, align="left", space_after=4)
    # Figure legends then figures
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, "Figures", size=12, bold=True, align="left", space_after=8)
    for fn, cap in MAIN_FIGS_NEW:
        add_figure(doc, fn, cap)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, "Tables", size=12, bold=True, align="left", space_after=8)
    for fn, cap in MAIN_TABLES_NEW:
        add_table(doc, pd.read_csv(fn), cap, fontsize=8)
    doc.save(path)
    return body, path


def build_sup(path="Cementum_CRISPR_supplementary.docx"):
    doc = Document()
    _style(doc)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1.0)
    para(doc, "Supplementary Information", size=14, bold=True, align="center",
         space_after=6)
    para(doc, M.TITLE, size=12, italic=True, align="center", space_after=14)
    para(doc, "Supplementary Figures", size=12, bold=True, align="left", space_after=8)
    for fn, cap in SUP_FIGS_NEW:
        add_figure(doc, fn, cap)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    para(doc, "Supplementary Tables", size=12, bold=True, align="left", space_after=8)
    for fn, cap in SUP_TABLES_NEW:
        df = pd.read_csv(fn)
        add_table(doc, df, cap, max_rows=200, fontsize=8)
    doc.save(path)
    return path
